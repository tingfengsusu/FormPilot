"""导入：Excel / CSV → 落库 → 逐条校验（设计文档 5.1）。

关键点（附录 A.4「导入前必须补的三件事」在此落地）：

1. **日期规则**：保险期限列常常没有年份（``8.22-8.21``），统一按
   「学平险 1 年期、止日期 = 次年」补年份；缺区间/非法日期/时长异常 → 异常
2. **学校字典 / 年级字典**：导入时归一化，匹配不上**不猜**，直接进「数据异常」
3. **占位符**：``（投保人姓名）`` 这类提示文字被 OCR 识别进来的，当空处理

导入即校验，逐条写 ``validation`` 表；全部通过 → 「未确认」，
有任何一项不通过 → 「数据异常」（带原因）。
"""
from __future__ import annotations

import csv
import os
from dataclasses import dataclass, field
from typing import (Any, Dict, Iterable, List, Mapping, Optional, Sequence,
                    Tuple)

from .db import BatchDB
from .dictionaries import MATCH_AMBIGUOUS, MATCH_UNMATCHED, DictionarySet
from .models import (STATUS_DATA_ERROR, STATUS_PENDING, Applicant, Batch,
                     ValidationResult)
from .recipe import Recipe
from .validator import (Validator, apply_period_floor, derive_student_level,
                        normalize_phone, normalize_text, resolve_period,
                        resolve_period_floor)


@dataclass
class ImportReport:
    """一次导入的结果（导入报告：哪些放行、哪些异常）。"""

    batch: Optional[Batch] = None
    total_rows: int = 0
    imported: int = 0
    passed: int = 0
    data_error: int = 0
    skipped: int = 0
    #: 字典未命中的原始值：{"school": ["清花园", ...], ...}
    unmatched: Dict[str, List[str]] = field(default_factory=dict)
    #: 逐条异常摘要：[(applicant_id, 原因)]
    errors: List[Tuple[int, str]] = field(default_factory=list)
    #: 行级硬错误（整行读不进来）：[(行号, 原因)]
    row_errors: List[Tuple[int, str]] = field(default_factory=list)
    #: 字典权威值是否已由页面探针产出
    options_ready: bool = True
    #: 尚未探针产出选项的字典名（school / grade / class_name）
    missing_options: List[str] = field(default_factory=list)
    #: 起期被按「页面只允许今天及之后」调整的条数
    period_adjusted: int = 0
    #: 调整样例（供人工确认）：["学生甲：2026-08-22 → 2026-09-19（止期同步顺延）"]
    period_adjust_samples: List[str] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)

    @property
    def summary(self) -> str:
        lines = [
            f"批次：{self.batch.name if self.batch else '-'}"
            f"（id={self.batch.id if self.batch else '-'}）",
            f"读取 {self.total_rows} 行，导入 {self.imported} 条"
            f"（跳过 {self.skipped} 行）",
            f"校验通过 {self.passed} 条 → 未确认；"
            f"数据异常 {self.data_error} 条 → 待人工修正",
        ]
        if self.row_errors:
            lines.append(f"行级错误 {len(self.row_errors)} 行：")
            lines += [f"  · 第 {r} 行：{msg}" for r, msg in self.row_errors[:10]]
        for key, values in self.unmatched.items():
            if values:
                lines.append(f"字典未命中（{key}）{len(values)} 种："
                             f"{'、'.join(values[:10])}"
                             + ("…" if len(values) > 10 else ""))
        if self.period_adjusted:
            lines.append(
                f"⚠ 保险期间起期已调整 {self.period_adjusted} 条"
                f"（页面只允许今天及之后的起期，台账原起期更早；止期同步顺延、"
                f"保持原时长）——请人工确认：")
            lines += [f"  · {sample}" for sample in self.period_adjust_samples[:5]]
            if self.period_adjusted > len(self.period_adjust_samples[:5]):
                lines.append(f"  · …另有 {self.period_adjusted - 5} 条，"
                             f"逐条明细见 validation 表（rule=period_adjusted）")
        if not self.options_ready:
            lines.append(
                f"⚠ 页面下拉选项尚未探针产出（{'、'.join(self.missing_options) or '全部'}）："
                f"无法确认这些值能否被选中，已全部计入「数据异常」"
                f"——请先运行 tools/probe_options.py 抓取选项后重新导入"
                f"（这是刻意的严格行为：不猜、不放行，避免到录入阶段才大面积失败）")
        return "\n".join(lines)


# ----------------------------------------------------------------------
# 读取文件
# ----------------------------------------------------------------------
def read_table(path: str, sheet: Optional[str] = None) -> Tuple[List[str], List[dict]]:
    """读取 Excel / CSV 的表头与数据行。

    :return: (headers, rows)，rows 为 ``{列名: 值}`` 列表
    """
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_excel(path, sheet)
    if ext == ".xls":
        raise ValueError("不支持旧版 .xls，请另存为 .xlsx 或 .csv")
    return _read_csv(path)


def _read_excel(path: str, sheet: Optional[str] = None) -> Tuple[List[str], List[dict]]:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover - 取决于环境
        raise RuntimeError("读取 Excel 需要 openpyxl：pip install openpyxl") from e

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows_iter = ws.iter_rows(values_only=True)

    headers: List[str] = []
    header_row = 0
    for row_no, row in enumerate(rows_iter, start=1):
        if row and any(str(c).strip() for c in row if c is not None):
            headers = ["" if c is None else str(c).strip() for c in row]
            header_row = row_no
            break

    sheet_name = getattr(ws, "title", "") or ""
    rows: List[dict] = []
    for row_no, row in enumerate(rows_iter, start=header_row + 1):
        if row is None or not any(c is not None and str(c).strip() for c in row):
            continue  # 全空行（**不重排行号**：__row__ 就是 Excel 里人看到的行号）
        data = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
        data["__row__"] = row_no
        data["__sheet__"] = sheet_name
        rows.append(data)
    wb.close()
    return headers, rows


def _read_csv(path: str) -> Tuple[List[str], List[dict]]:
    """读 CSV（utf-8-sig 优先，失败回退 GBK）。"""
    last_error: Optional[Exception] = None
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            with open(path, "r", encoding=encoding, newline="") as f:
                reader = csv.DictReader(f)
                headers = [h.strip() if h else "" for h in (reader.fieldnames or [])]
                rows = []
                for row_no, row in enumerate(reader, start=2):   # 第 1 行是表头
                    data = {k.strip() if k else k: v for k, v in row.items()}
                    data["__row__"] = row_no
                    rows.append(data)
            return headers, rows
        except UnicodeDecodeError as e:
            last_error = e
            continue
    raise RuntimeError(f"CSV 解码失败：{last_error}")


# ----------------------------------------------------------------------
# 列名映射
# ----------------------------------------------------------------------
class ColumnMapper:
    """把源表列名映射到明细字段名（支持多别名 + 占位符识别）。"""

    def __init__(self, mapping_config: dict):
        self.columns: Dict[str, List[str]] = {
            str(k): [str(a) for a in (v or [])]
            for k, v in (mapping_config.get("columns") or {}).items()
        }
        self.ignore_columns = {normalize_text(c) for c in
                               (mapping_config.get("ignore_columns") or [])}
        self.placeholders = {normalize_text(v) for v in
                             (mapping_config.get("ignore_placeholder_values") or [])}
        #: 归一化别名 → 字段名（含原始别名长度，用于 contains 匹配取最长）
        self._alias_index: Dict[str, Tuple[str, str]] = {}
        for fieldname, aliases in self.columns.items():
            for alias in aliases:
                self._alias_index[normalize_text(alias)] = (fieldname, alias)

    def resolve(self, headers: Sequence[str]) -> Dict[str, str]:
        """返回 ``{字段名: 源表列名}``。

        三轮匹配：精确 → 归一化精确 → 包含（取最长别名，避免
        「缴款人姓名」被短别名「姓名」抢走）。
        """
        result: Dict[str, str] = {}
        used: set = set()

        for header in headers:
            if normalize_text(header) in self.ignore_columns:
                used.add(header)
        for fieldname, aliases in self.columns.items():
            for alias in aliases:
                if alias in headers:
                    result[fieldname] = alias
                    used.add(alias)
                    break

        norm_headers = {normalize_text(h): h for h in headers if h not in used}
        for fieldname, aliases in self.columns.items():
            if fieldname in result:
                continue
            for alias in aliases:
                hit = norm_headers.get(normalize_text(alias))
                if hit is not None:
                    result[fieldname] = hit
                    used.add(hit)
                    break

        contains_candidates: List[Tuple[int, str, str]] = []  # (别名长度, 字段, 列名)
        for header in headers:
            if header in used:
                continue
            key = normalize_text(header)
            if not key:
                continue
            for alias_norm, (fieldname, alias) in self._alias_index.items():
                if alias_norm and alias_norm in key:
                    contains_candidates.append((len(alias_norm), fieldname, header))
        contains_candidates.sort(key=lambda item: -item[0])
        taken: set = set()
        for _, fieldname, header in contains_candidates:
            if fieldname in result or header in taken:
                continue
            result[fieldname] = header
            taken.add(header)
        return result

    def clean(self, value: Any) -> str:
        """清洗单元格值：占位符/占位文字 → 空。"""
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            value = int(value)  # Excel 把手机号读成浮点的情况
        text = str(value).strip()
        if normalize_text(text) in self.placeholders:
            return ""
        return text


# ----------------------------------------------------------------------
# 导入器
# ----------------------------------------------------------------------
class Importer:
    """Excel/CSV → SQLite（含归一化与校验）。"""

    def __init__(self, db: BatchDB, recipe: Optional[Recipe] = None,
                 dictionaries: Optional[DictionarySet] = None,
                 validator: Optional[Validator] = None):
        self.db = db
        self.recipe = recipe or Recipe.load()
        self.dictionaries = dictionaries or self.recipe.build_dictionaries()
        self.validator = validator or self.recipe.build_validator(self.dictionaries)
        from .recipe import load_batch_settings
        # 基础字典 + import_mapping.local.yaml（「从样例表对齐字段」自动追加的别名）
        from core.column_dict import load_mapping

        mapping_cfg = load_mapping(self.recipe.mapping_path) or \
            load_batch_settings().get("import", {}).get("mapping", {})
        self.mapper = ColumnMapper(mapping_cfg if isinstance(mapping_cfg, dict) else {})
        self.validate_on_import = bool(
            (self.recipe.input_config.get("validate_on_import", True)))
        #: 起期最早允许值（实测：投保页「保险期间·起」只能选今天及之后）
        self.period_start_floor = resolve_period_floor(
            self.recipe.input_config.get("period_start_floor", ""))

    # ------------------------------------------------------------------
    def check_table(self, path: str) -> Tuple[bool, str]:
        """导入**前**预检：这份台账的表头能不能对得上当前品种。

        判定链：列名字典能认出哪些列 → 明细字段名 → recipe.property_map 反查 →
        与品种 field_mapping 的属性集取交集。一个都对不上就直接告诉用户，
        别等跑完一百条才发现全是错的。
        """
        try:
            headers, _rows = read_table(path)
        except Exception as e:  # noqa: BLE001 — 读不了也要给出可读原因
            return False, f"读表失败：{e}"
        if not headers:
            return False, "没读到表头"

        column_map = self.mapper.resolve(headers)          # {明细字段名: 源列名}
        inverse = {field: prop for prop, field in self.recipe.property_map.items()}
        mapped = {inverse[f] for f in column_map if f in inverse}

        product = None
        try:
            from core.product_library import ProductLibrary

            for item in ProductLibrary().list_products():
                if item.name == self.recipe.product_name:
                    product = item
                    break
        except Exception:  # noqa: BLE001
            product = None
        expected = {m.property for m in (product.field_mapping if product else [])}
        overlap = mapped & expected
        missing = expected - overlap
        preview = "、".join(str(h) for h in headers[:8]) + ("…" if len(headers) > 8 else "")

        # 判定分三种，**别都归成"品种对不上"**（实测：新业务的正确台账会被误杀）
        recognized = len(column_map)                      # 字典认出的列数
        header_total = max(1, len([h for h in headers if str(h).strip()]))
        rate = recognized / header_total

        if not overlap and rate < 0.5:
            # 表头**大部分都没被字典认出** = 多半是新业务的表 —— 别说"品种选错"，
            # 否则用户明明选对了品种却被告知"确认品种"（实测反馈）
            return False, (
                f"列名字典几乎认不出这份表的列名（只认出 {recognized}/{header_total} 列）：\n"
                f"· 表头：{preview}\n"
                "→ 若是**新业务**的表：先在品种编辑页用「📄 从样例表对齐字段」配一遍列名"
                "（别名会写进 import_mapping.local.yaml，之后预检才准）；\n"
                "→ 如果是选错了文件或表头不规范，请重新选择。")

        ratio = len(overlap) / len(expected) if expected else 0.0
        if ratio < 0.5:
            return False, (
                f"覆盖不足：这份台账只能提供 {len(overlap)}/{len(expected)} 个"
                f"当前品种需要的字段（{ratio:.0%}）：\n"
                f"· 能提供：{'、'.join(sorted(overlap)) or '（无）'}\n"
                f"· 缺少：{'、'.join(sorted(missing)) or '（无）'}\n"
                "→ 确认批量页顶部的「品种」是否选对；若是新业务，先对齐字段/补列名字典。")

        return True, (f"覆盖 {len(overlap)}/{len(expected)} 个字段"
                      f"（{ratio:.0%}）：{'、'.join(sorted(overlap))}")

    def import_file(self, path: str, batch_name: str = "",
                    sheet: Optional[str] = None) -> ImportReport:
        """导入一个 Excel/CSV 文件，返回导入报告。"""
        headers, rows = read_table(path, sheet=sheet)
        return self.import_rows(rows, headers, batch_name=batch_name,
                                source_file=path)

    def import_rows(self, rows: Sequence[dict], headers: Sequence[str],
                    batch_name: str = "", source_file: str = "") -> ImportReport:
        """导入内存中的行（测试与预览用）。"""
        report = ImportReport(options_ready=self.dictionaries.options_ready,
                              missing_options=self.dictionaries.missing_options())
        column_map = self.mapper.resolve(headers)
        missing = [c for c in ("name", "id_card") if c not in column_map]
        if missing:
            report.notes.append(
                f"⚠ 源表未找到关键列：{'、'.join(missing)}"
                f"（已识别列：{column_map}）")

        name = batch_name or _default_batch_name(source_file)
        report.batch = self.db.create_batch(name, source_file=source_file,
                                            product=self.recipe.product_name)

        applicants: List[Applicant] = []
        for index, row in enumerate(rows, start=2):  # 第 1 行是表头
            try:
                applicant = self._build_applicant(row, column_map)
            except Exception as e:  # noqa: BLE001 — 单行失败不影响整批
                report.row_errors.append((index, str(e)))
                report.skipped += 1
                continue
            if not any([applicant.name, applicant.id_card, applicant.phone]):
                report.skipped += 1  # 空行
                continue
            applicant.batch_id = report.batch.id
            applicant.id = self.db.add_applicant(applicant)
            applicants.append(applicant)

        report.total_rows = len(rows)
        report.imported = len(applicants)

        if self.validate_on_import:
            self._validate_all(applicants, report)

        self.db.update_batch_total(report.batch.id)
        report.batch = self.db.get_batch(report.batch.id)
        return report

    # ------------------------------------------------------------------
    def _build_applicant(self, row: dict, column_map: Dict[str, str]) -> Applicant:
        """一行源数据 → Applicant（含期间补年份、字典归一化、默认投保人）。"""
        get = lambda key: self.mapper.clean(row.get(column_map.get(key, "")))  # noqa: E731

        phone = normalize_phone(get("phone")) or get("phone")
        payer_phone = normalize_phone(get("payer_phone")) or get("payer_phone")

        period_text = get("period")
        period_start = ""
        period_end = ""
        if period_text:
            # 整列文本形态："8.22-8.21" → 按「止日期=次年」补年份
            period = resolve_period(period_text, ref_year=self.recipe.ref_year,
                                    period_days=self.validator.period_days)
            period_start, period_end = period.iso()
        if not period_start and get("period_start"):
            period_start = _to_iso(get("period_start"), self.recipe.ref_year)
        if not period_end and get("period_end"):
            period_end = _to_iso(get("period_end"), self.recipe.ref_year)

        adjust_notes: List[str] = []
        if period_start and period_end and self.period_start_floor is not None:
            period_start, period_end, note = self._apply_period_floor(
                period_start, period_end)
            if note:
                adjust_notes.append(note)

        school = self._canonical("school", get("school"))
        raw_grade = get("grade")
        grade = self._canonical("grade", raw_grade)
        class_name = self._canonical("class_name", get("class_name"))
        # 学生类别：由**原始**年级写法推导（归一化后的「1年级」分不出小学一年级还是大一）
        student_level = derive_student_level(raw_grade,
                                             self.recipe.student_level_rules)

        applicant_name = get("applicant_name")
        if not applicant_name:
            applicant_name = self._default_applicant_name()

        applicant = Applicant(
            batch_id=0,
            name=get("name"),
            id_card=get("id_card").upper(),
            phone=phone,
            payer_name=get("payer_name"),
            payer_id_card=get("payer_id_card").upper(),
            payer_phone=payer_phone,
            applicant_name=applicant_name,
            school=school,
            grade=grade,
            class_name=class_name,
            student_level=student_level,
            period_start=period_start,
            period_end=period_end,
            status=STATUS_PENDING,
            source_image=get("source_image"),
            source_row=int(row.get("__row__") or 0),
            source_sheet=str(row.get("__sheet__") or ""),
        )
        applicant.adjust_notes.extend(adjust_notes)
        return applicant

    def _apply_period_floor(self, period_start: str,
                            period_end: str) -> Tuple[str, str, str]:
        """起期早于页面允许的最早日期时整体平移（保持原时长），返回调整说明。

        策略与留痕见 :func:`batch.validator.apply_period_floor`：
        起期取今天、止期同步顺延，并逐条提醒人工确认——**不改源文件**。
        """
        start, end, note = apply_period_floor(
            period_start, period_end, self.period_start_floor)
        if note:
            note += (f"——台账原起期早于页面允许的最早日期 "
                     f"{self.period_start_floor.isoformat()}，请人工确认起保日")
        return start, end, note

    def _canonical(self, key: str, value: str) -> str:
        """字典归一化：命中则用权威值；未命中保留原值（校验时判异常）。"""
        if not value:
            return ""
        match = self.dictionaries.match(key, value)
        return match.canonical if match.ok else value

    def _default_applicant_name(self) -> str:
        """投保人留空 → 取当前操作者（附录 B.2）。"""
        return str(self.recipe.execute.get("current_operator", "") or "")

    def _validate_all(self, applicants: List[Applicant],
                      report: ImportReport) -> None:
        """逐条校验并落库 + 置状态（重复检测在批次内整体算一次）。"""
        duplicates = self._duplicate_ids(applicants)
        unmatched: Dict[str, List[str]] = {}

        for applicant in applicants:
            vreport = self.validator.validate(applicant, duplicate_ids=duplicates)
            # 起期调整留痕（不改变通过与否，只提醒人工确认；见 batch.yaml）
            if applicant.adjust_notes:
                report.period_adjusted += 1
                if len(report.period_adjust_samples) < 5:
                    short = applicant.adjust_notes[0].split("——")[0]
                    report.period_adjust_samples.append(
                        f"{applicant.name or '(未填姓名)'}：{short}")
                for note in applicant.adjust_notes:
                    vreport.results.insert(0, ValidationResult(
                        applicant_id=applicant.id or 0, field="period_start",
                        rule="period_adjusted", passed=True, detail=note))
            self.db.replace_validations(applicant.id, vreport.results)

            for result in vreport.results:
                if result.passed:
                    continue
                if result.rule == "dictionary":
                    raw = getattr(applicant, result.field, "")
                    bucket = unmatched.setdefault(result.field, [])
                    if raw and raw not in bucket:
                        bucket.append(raw)

            if vreport.passed:
                report.passed += 1
                continue
            report.data_error += 1
            reason = vreport.reason
            report.errors.append((applicant.id, reason))
            self.db.set_status(applicant.id, STATUS_DATA_ERROR,
                               operator="system", note="导入校验不通过",
                               fail_reason=reason)

        report.unmatched = unmatched

    @staticmethod
    def _duplicate_ids(applicants: Sequence[Applicant]) -> set:
        """同批次内「同一身份证号 + 同一保险期间」的 id 集合。"""
        groups: Dict[tuple, List[int]] = {}
        for applicant in applicants:
            key = (applicant.id_card, applicant.period_start, applicant.period_end)
            groups.setdefault(key, []).append(applicant.id)
        ids: set = set()
        for key, members in groups.items():
            if key[0] and len(members) > 1:
                ids.update(members)
        return ids


# ----------------------------------------------------------------------
# 辅助
# ----------------------------------------------------------------------
def _to_iso(text: str, ref_year: int) -> str:
    """把各种日期写法转成 ISO（缺年份用 ref_year）。"""
    from .validator import parse_date
    parsed = parse_date(text)
    if parsed is None:
        return ""
    if parsed.year == 1900:  # parse_date 的「无年份」哨兵
        return parsed.replace(year=ref_year).isoformat()
    return parsed.isoformat()


def _default_batch_name(path: str) -> str:
    """默认批次名：文件名 + 时间（可被调用方覆盖）。"""
    from .db import now_str
    base = os.path.splitext(os.path.basename(path))[0] if path else "手工导入"
    return f"{base} {now_str()[:10]}"


def map_row_to_properties(row: Mapping[str, Any], recipe: Any = None) -> dict:
    """把一行（表头→值）按列名字典 + ``property_map`` 映射成 ``{品种属性名: 值}``。

    拖文件到属性面板、单条「载入数据」都用这条链；``__row__`` 之类的内部键忽略。
    """
    from core.column_dict import load_mapping

    recipe = recipe or Recipe.load()
    mapper = ColumnMapper(load_mapping(recipe.mapping_path) or {})
    inverse = {field: prop for prop, field in recipe.property_map.items()}

    headers = [k for k in row.keys() if not str(k).startswith("__")]
    column_map = mapper.resolve(headers)              # {明细字段名: 源列名}
    values: dict = {}
    for fieldname, column in column_map.items():
        prop = inverse.get(fieldname)
        if not prop:
            continue
        value = str(row.get(column) or "").strip()
        if value:
            values[prop] = value
    return values


def _read_yaml(path: str) -> dict:
    import yaml
    if not path or not os.path.exists(path):
        return {}
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def build_importer(db_path: Optional[str] = None,
                   recipe: Optional[Recipe] = None) -> Importer:
    """便捷构造：按配置建库 + 建导入器。"""
    from .db import BatchDB as _DB
    recipe = recipe or Recipe.load()
    db = _DB(db_path or recipe.storage_path("db_path", "data/batch.db"))
    return Importer(db, recipe=recipe)


__all__ = ["Importer", "ImportReport", "ColumnMapper", "read_table",
           "build_importer", "MATCH_AMBIGUOUS", "MATCH_UNMATCHED"]
